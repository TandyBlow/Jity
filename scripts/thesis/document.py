"""
Generate FYP-1 Partial Thesis .docx
Context Engineering and Memory Management for Web-Based Chinese AI TRPG
"""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT = r"D:\others\2025-2026.2\fyp\PartialThesis_AI_TRPG_Context_Engineering.docx"

doc = Document()

# ── Page setup ──
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

# ── Style configuration ──
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
# Set East Asian font
rPr = style.element.get_or_add_rPr()
rFonts = OxmlElement('w:rFonts')
rFonts.set(qn('w:eastAsia'), '宋体')
rPr.insert(0, rFonts)

for level, size in [(1, 16), (2, 14), (3, 13)]:
    h_style = doc.styles[f'Heading {level}']
    h_font = h_style.font
    h_font.name = 'Times New Roman'
    h_font.size = Pt(size)
    h_font.bold = True
    h_font.color.rgb = RGBColor(0, 0, 0)
    h_style.paragraph_format.space_before = Pt(12)
    h_style.paragraph_format.space_after = Pt(6)
    h_rPr = h_style.element.get_or_add_rPr()
    h_rFonts = OxmlElement('w:rFonts')
    h_rFonts.set(qn('w:eastAsia'), '黑体')
    h_rPr.insert(0, h_rFonts)


def add_para(text, style_name='Normal', bold=False, alignment=None, font_size=None, spacing_after=None):
    """Add a paragraph with optional formatting."""
    p = doc.add_paragraph(style=style_name)
    run = p.add_run(text)
    if bold:
        run.bold = True
    if font_size:
        run.font.size = Pt(font_size)
    if alignment is not None:
        p.alignment = alignment
    if spacing_after is not None:
        p.paragraph_format.space_after = Pt(spacing_after)
    return p


def add_heading(text, level=1):
    """Add a heading."""
    h = doc.add_heading(text, level=level)
    return h


def add_page_break():
    doc.add_page_break()


